"""Download selected 2021 Census Profile characteristics for Vancouver DAs.

Run only AFTER the Vancouver DA extraction script has created Vancouver_DAs.gpkg.
Requires: requests, pyarrow, openpyxl (pip install requests pyarrow openpyxl)
"""

import csv
import io
import sqlite3
import time
from pathlib import Path

import pyarrow as pa
import pyarrow.parquet as pq
import requests
from openpyxl import Workbook, load_workbook
from openpyxl.cell import WriteOnlyCell
from openpyxl.styles import Font


# ============================================================
# SETTINGS — CHANGE THESE AS NEEDED
# ============================================================

PROJECT_FOLDER = Path("/Users/aidancarter/Desktop/CodexSandbox/VancouverDAs")
DA_FILE = PROJECT_FOLDER / "Vancouver_DAs.gpkg"  # Created by your second script
DA_LAYER = "Vancouver_DAs"

METADATA_FILE = Path(
    "/Users/aidancarter/Desktop/CodexSandbox/TorontoVoting/"
    "StatsCan/CensusProfile2021/Metadata/DF_DA_metadata.xlsx"
)

# Enter any Census Profile characteristic codes you want to download.
CHARACTERISTICS = [1474]  # Median value of dwellings ($)

OUTPUT_PARQUET = PROJECT_FOLDER / "Vancouver_DA_Census_download.parquet"
OUTPUT_EXCEL = PROJECT_FOLDER / "Vancouver_DA_Census_download.xlsx"

EXPECTED_DAS = 1018  # Change if your geographic selection changes
DA_BATCH_SIZE = 20
BASE_URL = "https://api.statcan.gc.ca/census-recensement/profile/sdmx/rest/data/STC_CP,DF_DA"


# ============================================================
# HELPERS
# ============================================================

def clean_code(value):
    if value is None:
        return None
    value = str(value).strip()
    if value.endswith(".0"):
        value = value[:-2]
    return value or None


def clean_value(value):
    if value is None or not str(value).strip():
        return None
    try:
        number = float(str(value).strip())
    except ValueError as exc:
        raise ValueError(f"Unexpected non-numeric OBS_VALUE: {value!r}") from exc
    return int(number) if number.is_integer() else number


def chunks(items, size):
    for start in range(0, len(items), size):
        yield items[start:start + size]


# ============================================================
# 1. READ DAUIDs FROM THE GEOPACKAGE (NO GEOPANDAS REQUIRED)
# ============================================================

if not DA_FILE.is_file():
    raise FileNotFoundError(f"Run your Vancouver DA extraction script first: {DA_FILE}")

print("Reading DAUIDs from:", DA_FILE)
with sqlite3.connect(DA_FILE) as conn:
    # A GeoPackage is a SQLite database. We only need the DAUID column.
    da_rows = conn.execute(f'SELECT DAUID FROM "{DA_LAYER}"').fetchall()

da_uids = [clean_code(row[0]) for row in da_rows]
if any(code is None or len(code) != 8 or not code.isdigit() for code in da_uids):
    raise ValueError("Invalid or missing 8-digit DAUID in GeoPackage")
if len(set(da_uids)) != len(da_uids):
    raise ValueError("Duplicate DAUIDs in GeoPackage")
da_uids = sorted(da_uids)
print("Vancouver DAs:", len(da_uids))
if len(da_uids) != EXPECTED_DAS:
    raise ValueError(f"Expected {EXPECTED_DAS} DAs, found {len(da_uids)}")

codes = [int(code) for code in CHARACTERISTICS]
if not codes or len(codes) != len(set(codes)):
    raise ValueError("Provide at least one characteristic code, with no duplicates")


# ============================================================
# 2. READ CHARACTERISTIC NAMES FROM STATSCAN METADATA
# ============================================================

if not METADATA_FILE.is_file():
    raise FileNotFoundError(f"Characteristic metadata not found: {METADATA_FILE}")

print("Reading characteristic metadata...")
wb = load_workbook(METADATA_FILE, read_only=True, data_only=True)
try:
    ws = wb["CL_CHARACTERISTIC"]
    wanted = {str(code) for code in codes}
    names = {}
    for row in ws.iter_rows(min_row=6, values_only=True):
        if len(row) > 1 and clean_code(row[0]) in wanted:
            names[int(clean_code(row[0]))] = str(row[1]).strip()
finally:
    wb.close()

if set(names) != set(codes):
    raise ValueError(f"Characteristic codes missing from metadata: {sorted(set(codes) - set(names))}")
for code in codes:
    print(f"  {code}: {names[code]}")


# ============================================================
# 3. DOWNLOAD DATA DIRECTLY FROM STATSCAN
# ============================================================

session = requests.Session()
session.headers.update({"Accept": "text/csv"})
values = {}
da_batches = list(chunks(da_uids, DA_BATCH_SIZE))
print(f"Downloading {len(codes)} characteristics for {len(da_uids)} DAs...")

for code in codes:
    for batch_number, batch in enumerate(da_batches, start=1):
        dguids = "+".join("2021S0512" + uid for uid in batch)
        url = f"{BASE_URL}/A5.{dguids}.1.{code}.1"
        print(f"Characteristic {code}: batch {batch_number}/{len(da_batches)}")

        for attempt in range(1, 5):
            try:
                response = session.get(url, timeout=120)
                response.raise_for_status()
                break
            except requests.RequestException as exc:
                if attempt == 4:
                    raise RuntimeError(f"Statistics Canada request failed: {url}") from exc
                print(f"  Attempt {attempt} failed: {exc}")
                time.sleep(attempt * 2)

        reader = csv.DictReader(io.StringIO(response.content.decode("utf-8-sig")))
        required = {"ALT_GEO_CODE", "CHARACTERISTIC", "OBS_VALUE"}
        if not required.issubset(reader.fieldnames or []):
            raise ValueError(f"Unexpected StatsCan response fields: {reader.fieldnames}")

        for row in reader:
            returned_code = clean_code(row["CHARACTERISTIC"])
            if returned_code != str(code):
                continue
            da_uid = clean_code(row["ALT_GEO_CODE"])
            if da_uid not in batch:
                raise ValueError(f"Unexpected DAUID in response: {da_uid}")
            key = (da_uid, code)
            if key in values:
                raise ValueError(f"Duplicate DAUID + characteristic in response: {key}")
            values[key] = clean_value(row["OBS_VALUE"])


# ============================================================
# 4. VERIFY EXACT COVERAGE AND BUILD NEW TABLE
# ============================================================

expected_keys = {(uid, code) for uid in da_uids for code in codes}
if set(values) != expected_keys:
    missing = sorted(expected_keys - set(values))
    extra = sorted(set(values) - expected_keys)
    raise ValueError(f"API results incomplete. Missing (first 10): {missing[:10]}; extra: {extra[:10]}")

sorted_keys = sorted(expected_keys)
records = {
    "DAUID": [uid for uid, _ in sorted_keys],
    "Characteristic": [code for _, code in sorted_keys],
    "Characteristic_Name": [names[code] for _, code in sorted_keys],
    "Value": [values[key] for key in sorted_keys],
}

table = pa.table({
    "DAUID": pa.array(records["DAUID"], type=pa.string()),
    "Characteristic": pa.array(records["Characteristic"], type=pa.int32()),
    "Characteristic_Name": pa.array(records["Characteristic_Name"], type=pa.string()),
    "Value": pa.array(records["Value"], type=pa.float64()),
})

print("\nValidation:")
print("  Unique DAUIDs:", len(da_uids))
print("  Characteristics:", len(codes))
print("  Rows:", table.num_rows)
print("  Non-null values:", sum(value is not None for value in records["Value"]))
print("  Null values:", sum(value is None for value in records["Value"]))
assert table.num_rows == len(da_uids) * len(codes)

# Do not overwrite files from previous runs.
for output in (OUTPUT_PARQUET, OUTPUT_EXCEL):
    if output.exists():
        raise FileExistsError(f"Output already exists. Rename or move it before running: {output}")


# ============================================================
# 5. WRITE AND VALIDATE PARQUET
# ============================================================

pq.write_table(table, OUTPUT_PARQUET, compression="snappy")
if not pq.read_table(OUTPUT_PARQUET).equals(table):
    raise ValueError("Parquet validation failed")
print("Parquet validated:", OUTPUT_PARQUET)


# ============================================================
# 6. WRITE AND VALIDATE EXCEL
# ============================================================

excel_wb = Workbook(write_only=True)
excel_ws = excel_wb.create_sheet("Census Profile")
excel_ws.freeze_panes = "A2"
excel_ws.auto_filter.ref = f"A1:D{table.num_rows + 1}"
for letter, width in {"A": 15, "B": 19, "C": 75, "D": 18}.items():
    excel_ws.column_dimensions[letter].width = width

headers = ["DAUID", "Characteristic", "Characteristic_Name", "Value"]
header_cells = []
for header in headers:
    cell = WriteOnlyCell(excel_ws, value=header)
    cell.font = Font(bold=True)
    header_cells.append(cell)
excel_ws.append(header_cells)

for uid, code in sorted_keys:
    uid_cell = WriteOnlyCell(excel_ws, value=uid)
    uid_cell.number_format = "@"
    excel_ws.append([uid_cell, code, names[code], values[(uid, code)]])
excel_wb.save(OUTPUT_EXCEL)

check_wb = load_workbook(OUTPUT_EXCEL, read_only=True, data_only=True)
try:
    check_ws = check_wb["Census Profile"]
    check_rows = check_ws.iter_rows(min_row=2, values_only=True)
    count = 0
    for expected_key, row in zip(sorted_keys, check_rows):
        uid, code = expected_key
        if (str(row[0]), int(row[1]), row[2]) != (uid, code, names[code]):
            raise ValueError(f"Excel record mismatch at row {count + 2}")
        observed = row[3]
        expected = values[expected_key]
        if observed != expected:
            raise ValueError(f"Excel value mismatch at row {count + 2}")
        count += 1
    if count != len(sorted_keys) or next(check_rows, None) is not None:
        raise ValueError("Excel row count mismatch")
finally:
    check_wb.close()

print("Excel validated:", OUTPUT_EXCEL)
print("\nCOMPLETE — Source GeoPackage unchanged.")
